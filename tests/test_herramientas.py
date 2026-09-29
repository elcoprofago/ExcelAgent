import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from openpyxl import Workbook, load_workbook

import excel
import herramientas as at


FILAS = (('Cliente', 'Importe', 'Zona'),
         ('Zenit', 700, 'Norte'),
         ('ACME', 1000, 'Norte'),
         ('Beta', 250, 'Sur'),
         ('ACME', 1000, 'Norte'))


@pytest.fixture(scope='module')
def sesion():
    s = excel.Sesion(excel.Bus())
    yield s
    s.salir(cerrar_excel=True, descartar=True)


class Pedidos:
    'Confirmacion falsa: registra lo que se pregunto y contesta lo que diga la prueba.'
    def __init__(self, respuesta=True):
        self.respuesta = respuesta
        self.vistos = []

    def __call__(self, kind, title, detail):
        self.vistos.append((kind, title, detail))
        return self.respuesta


@pytest.fixture
def libro(sesion, tmp_path):
    def abrir(filas=FILAS, respuesta=True, approval='ask', nombre='prueba.xlsx'):
        ruta = tmp_path / nombre
        wb = Workbook()
        ws = wb.active
        ws.title = 'Datos'
        for f in filas:
            ws.append(list(f))
        wb.create_sheet('Control')['A1'] = 'NO TOCAR'
        wb.save(str(ruta))
        sesion.abrir(str(ruta), visible=False, respaldo=False, minimizar=False)
        pedidos = Pedidos(respuesta)
        return at.ToolBox(sesion, pedidos, approval), pedidos, str(ruta)
    yield abrir
    sesion.cerrar(False)


def celda(s, ref, hoja='Datos'):
    return s.wb.Worksheets(hoja).Range(ref).Value2


def control_intacto(s):
    return celda(s, 'A1', 'Control') == 'NO TOCAR'


# --- protocolo

def test_sin_libro_abierto_devuelve_error_legible(sesion):
    tb = at.ToolBox(sesion)
    assert tb.execute('workbook_info', {}).startswith('ERROR: no workbook is open')


def test_herramienta_desconocida_y_argumentos_malos(libro):
    tb, _, _ = libro()
    assert tb.execute('borrar_todo', {}).startswith("ERROR: unknown tool 'borrar_todo'")
    assert tb.execute('read_range', {'rango': 'A1'}).startswith('ERROR: bad arguments for read_range')
    with pytest.raises(at.ToolError):
        at.parse_args('{"a": ')


def test_error_de_excel_vuelve_como_texto(libro):
    tb, _, _ = libro()
    r = tb.execute('read_range', {'range': 'ZZZZ1'})
    assert r.startswith('ERROR:') and 'valid range' in r


def test_specs_y_metodos_coinciden():
    for s in at.SPECS:
        assert hasattr(at.ToolBox, 'tool_' + s['function']['name'])
    assert at.STATE_CHANGING <= at.NAMES


# --- lectura

def test_workbook_info_muestra_hojas_y_datos(libro):
    tb, _, ruta = libro()
    r = tb.execute('workbook_info', {})
    assert ruta in r
    assert "Sheet 'Datos': used range A1:C5" in r
    assert 'Cliente | Importe | Zona' in r
    assert "Sheet 'Control'" in r


def test_read_range_valores_formulas_y_errores(libro, sesion):
    tb, _, _ = libro()
    tb.execute('write_range', {'start_cell': 'E1', 'values': [['=SUM(B2:B5)', '=NOEXISTE(1)']]})
    r = tb.execute('read_range', {'range': 'E1:F1'})
    assert '1 | 2950 | #NAME?' in r
    r = tb.execute('read_range', {'range': 'E1', 'formulas': True})
    assert '=SUM(B2:B5)' in r


def test_read_range_grande_se_recorta(libro):
    tb, _, _ = libro()
    r = tb.execute('read_range', {'range': 'A1:J1000'})
    assert '[only the first 200 rows x 10 columns' in r


def test_find_en_todas_las_hojas(libro):
    tb, _, _ = libro()
    r = tb.execute('find', {'text': 'acme'})
    assert "'Datos'!A3" in r and "'Datos'!A5" in r
    assert 'was not found' in tb.execute('find', {'text': 'inexistente'})


# --- escritura y confirmaciones

def test_write_range_en_celdas_vacias_no_pregunta(libro, sesion):
    tb, pedidos, _ = libro()
    r = tb.execute('write_range', {'start_cell': 'E2', 'values': [['Total', '=B2*2'], [1.5, None]]})
    assert not pedidos.vistos
    assert celda(sesion, 'F2') == 1400
    assert celda(sesion, 'E3') == 1.5
    assert 'Wrote 2 x 2' in r


def test_write_range_que_pisa_datos_pregunta_y_respeta_el_no(libro, sesion):
    tb, pedidos, _ = libro(respuesta=False)
    r = tb.execute('write_range', {'start_cell': 'A2', 'values': [['Pisado']]})
    assert r.startswith('DENIED')
    assert pedidos.vistos and pedidos.vistos[0][0] == 'edit'
    assert 'A2=Zenit' in pedidos.vistos[0][2]
    assert celda(sesion, 'A2') == 'Zenit'


def test_write_range_avisa_formula_con_error(libro):
    tb, _, _ = libro()
    r = tb.execute('write_range', {'start_cell': 'E1', 'values': [['=SUMA(B2:B5)']]})
    assert 'Cells with errors: E1 #NAME?' in r


def test_write_range_sin_celdas_es_un_error(libro, sesion):
    # Regresion (Qwen3.5-9B local): [[]] contestaba 'Wrote 1 x 0 cells' sin escribir nada.
    tb, pedidos, _ = libro()
    for vacio in ([[]], [[], []]):
        r = tb.execute('write_range', {'start_cell': 'B6', 'values': vacio})
        assert r.startswith('ERROR: values has no cells'), r
    assert sesion.hoja('Datos').Range('A6:B6').Value2 == ((None, None),)
    assert pedidos.vistos == []


def test_approval_edits_no_pregunta_ediciones_pero_si_borrados(libro, sesion):
    tb, pedidos, _ = libro(respuesta=False, approval='edits')
    tb.execute('write_range', {'start_cell': 'A2', 'values': [['Nuevo']]})
    assert celda(sesion, 'A2') == 'Nuevo'
    assert not pedidos.vistos
    r = tb.execute('delete_rows_columns', {'where': '3'})
    assert r.startswith('DENIED') and pedidos.vistos[0][0] == 'delete'
    assert celda(sesion, 'A3') == 'ACME'


def test_fill_formula_ajusta_las_referencias(libro, sesion):
    tb, _, _ = libro()
    tb.execute('fill_formula', {'range': 'D2:D5', 'formula': '=B2*2'})
    assert celda(sesion, 'D2') == 1400
    assert celda(sesion, 'D4') == 500
    assert sesion.wb.Worksheets('Datos').Range('D5').Formula == '=B5*2'


def test_format_range(libro, sesion):
    tb, _, _ = libro()
    r = tb.execute('format_range', {'range': 'A1:C1', 'bold': True, 'fill_color': 'amarillo', 'font_color': '#0070C0'})
    h = sesion.wb.Worksheets('Datos')
    assert h.Range('B1').Font.Bold
    assert h.Range('B1').Interior.Color == excel.color_bgr('#FFFF00')
    assert 'bold' in r
    tb.execute('format_range', {'range': 'B2:B5', 'number_format': 'integer', 'borders': True})
    # Se mide lo que se VE: leer NumberFormat sufre la misma traduccion que escribirlo y no detectaba el error.
    h.Range('B2').Value2 = 1234567
    h.Columns('B').ColumnWidth = 30     # .Text muestra ### si no entra
    assert h.Range('B2').Text == '1.234.567'
    h.Range('B3').Value2 = 1234.5
    tb.execute('format_range', {'range': 'B3', 'number_format': 'currency'})
    assert h.Range('B3').Text.replace(' ', '') == '$1.234,50'
    tb.execute('format_range', {'range': 'B3', 'number_format': '0.00%'})
    assert h.Range('B3').Text == '123450,00%'
    assert tb.execute('format_range', {'range': 'A1', 'fill_color': 'fucsia chillon'}).startswith('ERROR: unknown color')


def test_sort_por_encabezado_mantiene_las_filas_juntas(libro, sesion):
    tb, _, _ = libro()
    r = tb.execute('sort_range', {'key_column': 'Importe', 'descending': True})
    assert 'by column B' in r
    assert celda(sesion, 'A1') == 'Cliente'
    assert celda(sesion, 'B2') == 1000 and celda(sesion, 'A2') == 'ACME'
    assert celda(sesion, 'B5') == 250 and celda(sesion, 'A5') == 'Beta'
    assert control_intacto(sesion)


def test_remove_duplicates_pregunta_y_borra(libro, sesion):
    tb, pedidos, _ = libro()
    r = tb.execute('remove_duplicates', {})
    assert pedidos.vistos[0][0] == 'delete'
    assert 'Removed 1 duplicate' in r
    assert celda(sesion, 'A5') is None
    assert control_intacto(sesion)


def test_replace_cuenta_y_respeta_el_no(libro, sesion):
    tb, _, _ = libro(respuesta=False)
    assert tb.execute('replace', {'find': 'Norte', 'replace_with': 'N'}).startswith('DENIED')
    assert celda(sesion, 'C2') == 'Norte'
    tb.approval = 'all'
    r = tb.execute('replace', {'find': 'Norte', 'replace_with': 'N'})
    assert 'about 3 cell' in r
    assert celda(sesion, 'C2') == 'N' and celda(sesion, 'C4') == 'Sur'


def test_insertar_y_borrar_filas(libro, sesion):
    tb, _, _ = libro()
    tb.execute('insert_rows_columns', {'where': '2'})
    assert celda(sesion, 'A2') is None and celda(sesion, 'A3') == 'Zenit'
    tb.execute('delete_rows_columns', {'where': '2'})
    assert celda(sesion, 'A2') == 'Zenit'
    assert tb.execute('insert_rows_columns', {'where': 'A1:B2'}).startswith('ERROR')


def test_clear_range(libro, sesion):
    tb, pedidos, _ = libro()
    tb.execute('clear_range', {'range': 'C2:C5'})
    assert pedidos.vistos and celda(sesion, 'C2') is None and celda(sesion, 'B2') == 700


def test_hojas(libro, sesion):
    tb, pedidos, _ = libro()
    tb.execute('sheet', {'action': 'add', 'name': 'Resumen'})
    tb.execute('sheet', {'action': 'rename', 'name': 'resumen', 'new_name': 'Totales'})
    assert 'Totales' in sesion.hojas()
    pedidos.respuesta = False
    assert tb.execute('sheet', {'action': 'delete', 'name': 'Totales'}).startswith('DENIED')
    pedidos.respuesta = True
    tb.execute('sheet', {'action': 'delete', 'name': 'Totales'})
    assert sesion.hojas() == ['Datos', 'Control']


def test_grafico_formatos_condicionales_lista_e_inmovilizar(libro, sesion):
    tb, _, _ = libro()
    h = sesion.wb.Worksheets('Datos')
    assert 'chart' in tb.execute('add_chart', {'range': 'A1:B5', 'chart_type': 'column', 'title': 'Ventas'})
    assert h.ChartObjects().Count == 1
    for args in ({'range': 'B2:B5', 'rule': 'greater_than', 'value': 500},
                 {'range': 'B2:B5', 'rule': 'between', 'value': 100, 'value2': 800},
                 {'range': 'C2:C5', 'rule': 'text_contains', 'value': 'Sur', 'fill_color': 'verde'},
                 {'range': 'A2:A5', 'rule': 'duplicates'},
                 {'range': 'B2:B5', 'rule': 'data_bar'}):
        r = tb.execute('conditional_format', args)
        assert not r.startswith('ERROR'), (args, r)
    assert h.Range('B2:B5').FormatConditions.Count == 3
    r = tb.execute('data_validation_list', {'range': 'C2:C5', 'items': ['Norte', 'Sur', 'Este']})
    assert not r.startswith('ERROR'), r
    assert h.Range('C2').Validation.Type == 3
    sep = at.list_separator(sesion.app)
    assert h.Range('C2').Validation.Formula1.split(sep) == ['Norte', 'Sur', 'Este']
    r = tb.execute('freeze_panes', {'cell': 'A2', 'sheet': 'Datos'})
    assert not r.startswith('ERROR'), r
    assert sesion.wb.Windows(1).FreezePanes


def test_tabla_y_autofiltro(libro, sesion):
    tb, _, _ = libro()
    tb.execute('autofilter', {})
    r = tb.execute('create_table', {})
    assert 'with 4 data rows' in r, r
    assert tb.execute('create_table', {}).startswith('ERROR')


# --- archivos: nunca se pisa nada sin permiso

def test_save_pregunta_y_el_no_deja_el_archivo_como_estaba(libro, sesion):
    tb, pedidos, ruta = libro(respuesta=False)
    tb.execute('write_range', {'start_cell': 'E1', 'values': [['nuevo']]})
    assert tb.execute('save', {}).startswith('DENIED')
    assert load_workbook(ruta)['Datos']['E1'].value is None
    pedidos.respuesta = True
    assert 'Saved' in tb.execute('save', {})
    assert load_workbook(ruta)['Datos']['E1'].value == 'nuevo'
    assert load_workbook(ruta)['Control']['A1'].value == 'NO TOCAR'


def test_save_copy_y_csv_nunca_pisan(libro, tmp_path):
    tb, _, ruta = libro()
    r = tb.execute('save_copy', {'name': 'Copia'})
    assert os.path.isfile(tmp_path / 'Copia.xlsx'), r
    assert tb.execute('save_copy', {'name': 'Copia'}).startswith('ERROR')
    existente = tmp_path / 'viejo.csv'
    existente.write_text('CONTROL', encoding='utf-8')
    assert tb.execute('export_csv', {'path': str(existente)}).startswith('ERROR')
    assert existente.read_text(encoding='utf-8') == 'CONTROL'
    r = tb.execute('export_csv', {'path': str(tmp_path / 'nuevo')})
    texto = (tmp_path / 'nuevo.csv').read_text(encoding='utf-8-sig')
    assert texto.splitlines()[1] == 'Zenit;700;Norte', r


# --- adjuntos

def test_adjunto_leer_y_pegar(libro, sesion, tmp_path):
    tb, _, _ = libro()
    assert tb.execute('read_attachment', {}).startswith('ERROR: no attachments')
    archivo = tmp_path / 'lista.txt'
    archivo.write_text('Nombre;Codigo\nAna;007\nBeto;12\n', encoding='utf-8')
    sesion.adjuntos = [excel.procesar_adjunto(str(archivo))]
    try:
        assert 'lista.txt' in tb.execute('read_attachment', {})
        tb.execute('paste_attachment', {'start_cell': 'A1', 'sheet': 'Control'})
        h = sesion.wb.Worksheets('Control')
        assert h.Range('B2').Value2 == '007'     # los ceros a la izquierda se respetan
        assert h.Range('B3').Value2 == 12
    finally:
        sesion.adjuntos = []


def test_outlook_contacts_pasa_el_pedido_de_todo_el_buzon(libro, monkeypatch):
    # Regresion: el aviso de 'se miraron los 1200 mas recientes' ofrecia revisar el buzon entero, pero la herramienta
    # no tenia como pedirlo. Outlook falso con la carpeta de Contactos vacia: tiene que caer en los correos.
    import win32com.client

    class Vacio:
        def GetNamespace(self, _):
            return self

        def GetDefaultFolder(self, _):
            return self

        Items = ()

    llamados = []
    monkeypatch.setattr(win32com.client, 'GetActiveObject', lambda _: Vacio())
    monkeypatch.setattr(excel, 'op_contactos_correo', lambda s, todos=False, con_orden=False: llamados.append(todos) or 'ok')
    tb, _, _ = libro()
    tb.execute('outlook_contacts', {})
    tb.execute('outlook_contacts', {'whole_mailbox': True})
    assert llamados == [False, True]
