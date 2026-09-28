import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

import app

carpeta = tempfile.mkdtemp(prefix='excelagent_')
libro = os.path.join(carpeta, 'prueba.xlsx')

wb = Workbook()
ws = wb.active
ws.title = 'Ventas'
ws.append(['Producto', 'Cantidad', 'Precio'])
ws.append(['Detergente', 10, 1500.5])
ws.append(['Lavandina', 4, 800])
ws.append(['Jabon', 7, 325.25])
wb.save(libro)
print('libro de prueba:' + libro)


bus = app.Bus(sink=lambda ev: print('  ' + ev.hora + ' [' + ev.tipo + '] ' + ev.texto))
agente = app.Agente(bus)
agente.abrir(libro, visible=False, respaldo=True)

consignas = [
      'resumen'
    , 'escribir 99 en E1'
    , 'negrita en A1:C1'
    , 'sumar columna C'
    , 'crear hoja Nueva'
    , 'activar hoja Ventas'
    , 'ordenar por columna B descendente'
    , 'exportar csv ' + os.path.join(carpeta, 'salida.csv')
    , 'guardar'
]
for c in consignas:
    print('CONSIGNA: ' + c)
    agente.pedir(c)


agente.salir(cerrar_excel=True)

ver = load_workbook(libro)
h = ver['Ventas']
print('E1 = ' + str(h['E1'].value))
print('A1 negrita = ' + str(h['A1'].font.bold))
print('C5 total = ' + str(h['C5'].value))
print('hojas = ' + str(ver.sheetnames))
print('B2 tras ordenar desc = ' + str(h['B2'].value))
print('B4 tras ordenar desc = ' + str(h['B4'].value))
print('csv = ' + str(os.path.isfile(os.path.join(carpeta, 'salida.csv'))))
print('respaldos = ' + str([n for n in os.listdir(os.path.join(carpeta, '_excelagent_respaldos'))]))

