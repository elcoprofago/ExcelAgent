import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

import app


def test_varias_consignas_en_un_solo_mensaje(tmp_path):
    libro = tmp_path / 'multi.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Datos'
    ws.append(['Nombre', 'Importe'])
    ws.append(['Ana', 100])
    ws.append(['Beto', 250])
    wb.save(str(libro))

    bus = app.Bus()
    agente = app.Agente(bus)
    try:
        agente.abrir(str(libro), visible=False, respaldo=False)
        mensaje = chr(10).join(['escribir 111 en D1', 'negrita en A1:B1', 'sumar columna B'])
        respuesta = agente.pedir(mensaje)
        assert '111' in respuesta
        assert 'TOTAL' in respuesta
        agente.pedir('guardar')
    finally:
        agente.salir(cerrar_excel=True)

        ver = load_workbook(str(libro))
        h = ver['Datos']
        assert h['D1'].value == 111
        assert h['A1'].font.bold is True
        total = None
        for fila in range(1, h.max_row + 1):
            if h.cell(row=fila, column=1).value == 'TOTAL':
                total = h.cell(row=fila, column=2).value
        assert total is not None and str(total).startswith('=SUM')


def test_consigna_vacia_no_rompe(tmp_path):
    bus = app.Bus()
    agente = app.Agente(bus)
    assert 'No escribiste' in agente.pedir('   ')
