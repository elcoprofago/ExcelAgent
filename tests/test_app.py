import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

import excel
import herramientas

# Utilidades de excel.py que siguen en uso con el asesor. Las pruebas del interprete por reglas
# (interpretar, normalizar, ayuda) se fueron con el: ahora interpreta el modelo.


def test_valor_numeros_y_texto():
    assert excel.valor('42') == 42
    assert excel.valor('1500,50') == 1500.5
    assert excel.valor('texto libre') == 'texto libre'


def test_color_bgr():
    assert excel.color_bgr('#FF0000') == 255
    assert excel.color_bgr('#00FF00') == 65280


def test_col_num():
    assert excel.col_num('A') == 1
    assert excel.col_num('AA') == 27


def test_bus_guarda_eventos():
    bus = excel.Bus()
    bus.log('hola', 'ok')
    bus.progreso(50, 'mitad')
    assert len(bus.historial) == 2
    assert bus.historial[1].pct == 50


def test_partir_filas_y_valor_celda():
    assert excel.partir_filas('a,b' + chr(10) + 'c,d') == [['a', 'b'], ['c', 'd']]
    assert excel.partir_filas('x' + chr(9) + 'y') == [['x', 'y']]
    assert excel.valor_celda('1500') == 1500
    assert excel.valor_celda('011') == '011'
    assert excel.valor_celda('ACME') == 'ACME'


def test_procesar_adjunto_texto(tmp_path):
    ruta = tmp_path / 'contactos.csv'
    ruta.write_text('Nombre;Correo' + chr(10) + 'Ana;ana@x.com' + chr(10) + 'Beto;beto@x.com' + chr(10), encoding='utf-8')
    info = excel.procesar_adjunto(str(ruta))
    assert info['clase'] == 'texto'
    assert info['columnas'] == 2
    assert len(info['filas']) == 3
    assert 'Ana' in info['texto']


def test_ocr_lee_una_imagen(tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    ruta = tmp_path / 'contactos.png'
    img = Image.new('RGB', (760, 220), 'white')
    d = ImageDraw.Draw(img)
    fuente = ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf', 32)
    d.text((20, 20), 'CONTACTOS', fill='black', font=fuente)
    d.text((20, 80), 'Ana Perez  ana@x.com', fill='black', font=fuente)
    d.text((20, 140), 'Beto Diaz  beto@x.com', fill='black', font=fuente)
    img.save(ruta)
    info = excel.procesar_adjunto(str(ruta))
    assert info['clase'] == 'imagen'
    assert 'CONTACTOS' in info['texto'].upper()
    assert 'ana@x.com' in info['texto'].lower()


def test_con_limite_devuelve_el_valor():
    assert excel.con_limite(lambda: 42, 5) == 42


def test_con_limite_falla_si_no_responde():
    import time as reloj
    try:
        excel.con_limite(lambda: reloj.sleep(4), 1)
        raise AssertionError('deberia haber fallado por tiempo')
    except excel.ExcelError as exc:
        assert 'no respondio' in str(exc)


def test_con_limite_propaga_el_error_real():
    def rompe():
        raise ValueError('averia de prueba')
    try:
        excel.con_limite(rompe, 5)
        raise AssertionError('deberia haber propagado el error')
    except ValueError as exc:
        assert 'averia de prueba' in str(exc)


def test_limpiar_nombre():
    comilla = chr(39)
    dos = chr(34)
    assert excel.limpiar_nombre('Mauro Blanco' + comilla) == 'Mauro Blanco'
    assert excel.limpiar_nombre(dos + ': Ana Perez' + comilla + dos) == 'Ana Perez'
    assert excel.limpiar_nombre('  Beto  ') == 'Beto'


def test_correo_valido_descarta_basura():
    # Regresion: undisclosed-recipients y otros valores rotos llegaban a la planilla.
    assert excel.correo_valido(str()) == str()
    assert excel.correo_valido('undisclosed-recipients:') == str()
    assert excel.correo_valido('ana perez@x.com') == str()
    assert excel.correo_valido('ana@@x.com') == str()
    assert excel.correo_valido('ana@x.com') == 'ana@x.com'


def test_detener_corta_el_recorrido_de_correos():
    # El boton Detener debe cortar el trabajo en el proximo control.
    s = excel.Sesion(excel.Bus())
    import pythoncom
    pythoncom.CoInitialize()
    s.detener.set()
    try:
        excel.op_contactos_correo(s, False, False)
        raise AssertionError('deberia haber cortado el recorrido')
    except excel.ExcelError as exc:
        assert 'detenido' in str(exc)


# --- Salir nunca descarta cambios sin guardar ---

def libro(tmp_path, filas, nombre='libro.xlsx'):
    ruta = tmp_path / nombre
    wb = Workbook()
    ws = wb.active
    ws.title = 'Datos'
    for fila in filas:
        ws.append(list(fila))
    wb.save(str(ruta))
    return str(ruta)


def cerrar_todo(xl):
    # Limpieza de las pruebas: cierra sin guardar lo que haya quedado abierto.
    if xl is None:
        return
    try:
        while xl.Workbooks.Count:
            xl.Workbooks(1).Close(False)
        xl.Quit()
    except Exception:
        pass


def test_salir_con_cambios_pendientes_deja_el_libro_abierto(tmp_path):
    ruta = libro(tmp_path, [('original',)])
    s = excel.Sesion(excel.Bus())
    xl = None
    try:
        s.abrir(ruta, visible=False, respaldo=False, minimizar=False)
        xl = s.app
        herramientas.ToolBox(s, approval='all').tool_write_range('A1', [['modificado']])
        s.salir(cerrar_excel=True)
        assert xl.Workbooks.Count == 1
        assert xl.Workbooks(1).Worksheets(1).Range('A1').Value2 == 'modificado'
        assert xl.Visible
        assert xl.UserControl
    finally:
        cerrar_todo(xl)
    assert load_workbook(ruta).active['A1'].value == 'original'


def test_salir_sin_cambios_cierra_normalmente(tmp_path):
    ruta = libro(tmp_path, [('original',)])
    s = excel.Sesion(excel.Bus())
    xl = None
    try:
        s.abrir(ruta, visible=False, respaldo=False, minimizar=False)
        xl = s.app
        s.salir(cerrar_excel=False)
        assert xl.Workbooks.Count == 0
    finally:
        cerrar_todo(xl)
