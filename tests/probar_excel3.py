import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

import app

XL_MIN = -4140

carpeta = tempfile.mkdtemp(prefix='excelagent3_')
libro = os.path.join(carpeta, 'prueba3.xlsx')

wb = Workbook()
ws = wb.active
ws.title = 'Datos'
ws.append(['Nombre', 'Importe'])
ws.append(['Ana', 100])
ws.append(['Beto', 250])
wb.save(libro)
print('LIBRO=' + libro)

texto_adj = os.path.join(carpeta, 'contactos.txt')
with open(texto_adj, 'w', encoding='utf-8') as f:
    f.write('Nombre;Correo;Telefono' + chr(10))
    f.write('Ana Perez;ana@x.com;1145550001' + chr(10))
    f.write('Beto Diaz;beto@x.com;1145550002' + chr(10))
    f.write('Caro Gil;caro@x.com;1145550003' + chr(10))
print('ADJUNTO_TEXTO=' + texto_adj)

from PIL import Image, ImageDraw, ImageFont
img_adj = os.path.join(carpeta, 'precios.png')
img = Image.new('RGB', (900, 260), 'white')
d = ImageDraw.Draw(img)
fuente = ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf', 30)
y = 20
for linea in ('LISTA DE PRECIOS', 'Detergente 1500', 'Lavandina 800', 'TOTAL 2300'):
    d.text((20, y), linea, fill='black', font=fuente)
    y += 55
img.save(img_adj)
print('ADJUNTO_IMAGEN=' + img_adj)

bus = app.Bus(sink=lambda ev: print('  ' + ev.hora + ' [' + ev.nivel + '] ' + ev.texto))
agente = app.Agente(bus)
agente.abrir(libro, visible=True, respaldo=True, minimizar=True)
estado = agente.sesion.app.WindowState
print('ESTADO_VENTANA=' + str(estado) + ' (esperado ' + str(XL_MIN) + ')')
print('VISIBLE=' + str(agente.sesion.app.Visible))

print('--- adjunto de texto ---')
agente.adjuntar(texto_adj)
agente.pedir('analizar el adjunto')
agente.pedir('volcar el adjunto en A1')

print('--- adjunto de imagen con OCR ---')
agente.adjuntar(img_adj)
agente.pedir('analizar el adjunto')
agente.pedir('volcar el adjunto en A6')

print('--- crear tabla ---')
agente.pedir('crear una tabla en A1:C4')

print('--- contactos de Outlook ---')
agente.pedir('En la hoja designada crea una tabla que contenga un listado de contactos que surjan de Outlook')

print('--- mostrar excel y guardar ---')
agente.pedir('mostrar excel')
agente.pedir('guardar')
agente.salir(cerrar_excel=True)

ver = load_workbook(libro)
h = ver['Datos']
print('--- VERIFICACION con openpyxl ---')
print('A1=' + str(h['A1'].value) + ' | B1=' + str(h['B1'].value) + ' | C1=' + str(h['C1'].value))
print('A2=' + str(h['A2'].value) + ' | C2=' + str(h['C2'].value))
print('A6=' + str(h['A6'].value))
print('tablas=' + str(list(h.tables)))
print('filas=' + str(h.max_row))
