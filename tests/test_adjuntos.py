# Lectura de los formatos que acepta el boton '+' (adjuntos.py). Los archivos se arman aca mismo: .docx, .odt y .ods
# son zip con un XML minimo; .xls y .doc se guardan con Excel y Word en instancias propias.
import datetime
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from openpyxl import Workbook

import adjuntos
from excel import ExcelError

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
ODF = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
       'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
       'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')


def zip_con(ruta, nombre, xml):
    with zipfile.ZipFile(ruta, 'w') as z:
        z.writestr(nombre, xml)
    return str(ruta)


def celda_w(t):
    return f'<w:tc><w:p><w:r><w:t>{t}</w:t></w:r></w:p></w:tc>'


def test_docx_separa_las_tablas_del_texto(tmp_path):
    xml = (f'<w:document {W}><w:body><w:p><w:r><w:t>Lista de precios</w:t></w:r></w:p>'
           f'<w:tbl><w:tr>{celda_w("Producto")}{celda_w("Precio")}</w:tr>'
           f'<w:tr>{celda_w("Yerba")}{celda_w("3500")}</w:tr></w:tbl>'
           '<w:p><w:r><w:t>Fin</w:t></w:r></w:p></w:body></w:document>')
    r = adjuntos.leer(zip_con(tmp_path / 'precios.docx', 'word/document.xml', xml))
    assert r['clase'] == 'documento'
    assert [p['nombre'] for p in r['partes']] == ['Tabla 1']
    assert r['partes'][0]['filas'] == [['Producto', 'Precio'], ['Yerba', '3500']]
    assert 'Lista de precios' in r['texto'] and 'Fin' in r['texto']


def test_ods_con_tipos_y_sin_el_relleno_de_celdas_vacias(tmp_path):
    # LibreOffice guarda las filas y columnas vacias como 'repetidas' hasta el final de la hoja: no deben volverse un
    # millon de filas.
    fila = ('<table:table-row><table:table-cell office:value-type="string"><text:p>Pan</text:p></table:table-cell>'
            '<table:table-cell office:value-type="float" office:value="3"><text:p>3</text:p></table:table-cell>'
            '<table:table-cell office:value-type="date" office:date-value="2026-09-29"><text:p>29/09/26</text:p>'
            '</table:table-cell><table:table-cell table:number-columns-repeated="1000"/></table:table-row>')
    xml = (f'<office:document-content {ODF}><office:body><office:spreadsheet><table:table table:name="Compras">'
           '<table:table-row><table:table-cell office:value-type="string"><text:p>Item</text:p></table:table-cell>'
           '</table:table-row>' + fila +
           '<table:table-row table:number-rows-repeated="1048000"><table:table-cell table:number-columns-repeated="1024"/>'
           '</table:table-row></table:table></office:spreadsheet></office:body></office:document-content>')
    r = adjuntos.leer(zip_con(tmp_path / 'compras.ods', 'content.xml', xml))
    assert r['clase'] == 'planilla'
    p = r['partes'][0]
    assert p['nombre'] == 'Compras'
    assert p['filas'] == [['Item'], ['Pan', 3, datetime.date(2026, 9, 29)]]


def test_odt_lee_sus_tablas(tmp_path):
    xml = (f'<office:document-content {ODF}><office:body><office:text><text:p>Hola</text:p>'
           '<table:table table:name="Tabla1"><table:table-row><table:table-cell><text:p>A</text:p></table:table-cell>'
           '<table:table-cell><text:p>B</text:p></table:table-cell></table:table-row></table:table>'
           '</office:text></office:body></office:document-content>')
    r = adjuntos.leer(zip_con(tmp_path / 'nota.odt', 'content.xml', xml))
    assert r['clase'] == 'documento'
    assert r['partes'][0]['filas'] == [['A', 'B']]


def test_xlsx_todas_las_hojas_con_tipos(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ventas'
    ws.append(['Fecha', 'Monto'])
    ws.append([datetime.date(2026, 9, 1), 1500.5])
    wb.create_sheet('Vacia')
    wb.create_sheet('Clientes').append(['Ana', '007'])
    wb.save(tmp_path / 'libro.xlsx')
    r = adjuntos.leer(str(tmp_path / 'libro.xlsx'))
    assert [p['nombre'] for p in r['partes']] == ['Ventas', 'Clientes']
    assert r['partes'][0]['filas'][1] == [datetime.datetime(2026, 9, 1), 1500.5]
    assert r['partes'][1]['filas'] == [['Ana', '007']]
    assert '== Hoja Clientes' in r['texto']


def test_markdown_tabla_sin_la_linea_de_guiones(tmp_path):
    # Control: un Markdown sin tablas se lee como texto, linea por linea.
    (tmp_path / 'notas.md').write_text('# Notas\n\n| Nombre | Edad |\n|---|--:|\n| Ana | 30 |\n\nfin\n', encoding='utf-8')
    (tmp_path / 'solo.md').write_text('# Titulo\nuna linea\n', encoding='utf-8')
    r = adjuntos.leer(str(tmp_path / 'notas.md'))
    assert r['partes'][0]['filas'] == [['Nombre', 'Edad'], ['Ana', '30']]
    r = adjuntos.leer(str(tmp_path / 'solo.md'))
    assert r['partes'][0]['filas'] == [['# Titulo'], ['una linea']]


def test_binario_desconocido_se_rechaza_diciendo_que_acepta(tmp_path):
    (tmp_path / 'raro.bin').write_bytes(bytes(range(256)) * 4)
    with pytest.raises(ExcelError, match=r'\.xlsx.*\.docx'):
        adjuntos.leer(str(tmp_path / 'raro.bin'))


def test_xls_viejo_se_lee_con_una_instancia_propia_de_excel(tmp_path):
    import win32com.client
    xl = win32com.client.DispatchEx('Excel.Application')
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Add()
        wb.Worksheets(1).Name = 'Datos'
        wb.Worksheets(1).Range('A1:B2').Value2 = (('Codigo', 'Monto'), ("'007", 12.5))
        wb.SaveAs(str(tmp_path / 'viejo.xls'), 56)
        wb.Close(False)
    finally:
        xl.Quit()
    r = adjuntos.leer(str(tmp_path / 'viejo.xls'))
    assert r['partes'][0]['nombre'] == 'Datos'
    assert r['partes'][0]['filas'] == [['Codigo', 'Monto'], ['007', 12.5]]


def test_pdf_con_texto_saca_las_columnas_de_la_tabla(tmp_path):
    # El PDF lo arma Word ('Guardar como PDF'), como los que recibe el usuario.
    import pywintypes
    import win32com.client
    try:
        pywintypes.IID('Word.Application')
    except pywintypes.com_error:
        pytest.skip('Word no esta instalado: no hay con que armar el PDF de prueba')
    wd = win32com.client.DispatchEx('Word.Application')
    try:
        d = wd.Documents.Add()
        d.Content.Text = 'Lista de precios'
        d.Content.InsertParagraphAfter()
        t = d.Tables.Add(d.Paragraphs(d.Paragraphs.Count).Range, 2, 3)
        for j, v in enumerate(('Yerba mate 1 kg', '1.500,50', '01/09/2026'), 1):
            t.Cell(2, j).Range.Text = v
        for j, v in enumerate(('Producto', 'Precio', 'Fecha'), 1):
            t.Cell(1, j).Range.Text = v
        d.SaveAs2(str(tmp_path / 'precios.pdf'), 17)
        d.Close(False)
    finally:
        wd.Quit()
    r = adjuntos.leer(str(tmp_path / 'precios.pdf'))
    assert r['clase'] == 'documento'
    assert r['partes'][0]['filas'] == [['Producto', 'Precio', 'Fecha'], ['Yerba mate 1 kg', '1.500,50', '01/09/2026']]
    assert 'Lista de precios' in r['texto']


def test_pdf_escaneado_se_lee_por_ocr(tmp_path):
    import excel
    if not excel.ruta_tesseract():
        pytest.skip('No hay tesseract.exe en esta PC')
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new('RGB', (1200, 300), 'white')
    ImageDraw.Draw(img).text((40, 100), 'Maria 1144443333', fill='black', font=ImageFont.truetype('arial.ttf', 48))
    img.save(tmp_path / 'escaneado.pdf')
    r = adjuntos.leer(str(tmp_path / 'escaneado.pdf'))
    assert 'Maria 1144443333' in r['texto']
    assert 'OCR' in r['texto']          # el aviso de revisar lo leido


def test_pdf_sin_texto_ni_imagen_pide_otra_via(tmp_path):
    import pypdf
    w = pypdf.PdfWriter()
    w.add_blank_page(200, 200)
    with open(tmp_path / 'vacio.pdf', 'wb') as f:
        w.write(f)
    with pytest.raises(ExcelError, match='captura'):
        adjuntos.leer(str(tmp_path / 'vacio.pdf'))


def test_doc_viejo_se_lee_con_word(tmp_path):
    import pywintypes
    import win32com.client
    try:
        pywintypes.IID('Word.Application')
    except pywintypes.com_error:
        pytest.skip('Word no esta instalado: .doc y .rtf no se pueden leer en esta PC')
    wd = win32com.client.DispatchEx('Word.Application')
    try:
        d = wd.Documents.Add()
        t = d.Tables.Add(d.Content, 1, 2)
        t.Cell(1, 1).Range.Text = 'Producto'
        t.Cell(1, 2).Range.Text = 'Precio'
        d.SaveAs2(str(tmp_path / 'viejo.doc'), 0)
        d.Close(False)
    finally:
        wd.Quit()
    r = adjuntos.leer(str(tmp_path / 'viejo.doc'))
    assert r['partes'][0]['filas'] == [['Producto', 'Precio']]
