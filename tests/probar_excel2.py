import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

import app

carpeta = tempfile.mkdtemp(prefix='excelagent2_')
libro = os.path.join(carpeta, 'prueba2.xlsx')

wb = Workbook()
ws = wb.active
ws.title = 'Datos'
ws.append(['Cliente', 'Importe', 'Zona'])
ws.append(['ACME', 1000, 'Norte'])
ws.append(['ACME', 250, 'Sur'])
ws.append(['Zenit', 700, 'Norte'])
ws.append(['ACME', 1000, 'Norte'])
wb.save(libro)

bus = app.Bus()
agente = app.Agente(bus)
agente.abrir(libro, visible=False, respaldo=True)
print('LIBRO:' + libro)

consignas = [
      'leer A1:C3'
    , 'resumen'
    , 'promedio columna B'
    , 'maximo columna B'
    , 'negrita en A1:C1'
    , 'color rojo en A1:C1'
    , 'fondo amarillo en A1:C1'
    , 'formato moneda en B2:B5'
    , 'buscar ACME'
    , 'reemplazar Sur por Centro'
    , 'quitar duplicados'
    , 'autofiltro'
    , 'autoajustar'
    , 'crear hoja Extra'
    , 'renombrar hoja Extra a Borrador'
    , 'activar hoja Datos'
    , 'guardar'
]
for c in consignas:
    print('CONSIGNA: ' + c)
    print('  -> ' + str(agente.pedir(c)).replace(chr(10), ' / '))

agente.salir(cerrar_excel=True)

ver = load_workbook(libro)
h = ver['Datos']
print('--- VERIFICACION CON openpyxl ---')
print('A1 negrita = ' + str(h['A1'].font.bold))
print('A1 color = ' + str(h['A1'].font.color.rgb))
print('A1 fondo = ' + str(h['A1'].fill.start_color.rgb))
print('B2 formato = ' + str(h['B2'].number_format))
print('C3 = ' + str(h['C3'].value))
print('filas = ' + str(h.max_row))
print('hojas = ' + str(ver.sheetnames))
print('B6 (promedio) = ' + str(h['B6'].value))

