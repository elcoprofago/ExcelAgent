# Lectura de los archivos que el usuario adjunta con el boton del clip. Cada formato se convierte en texto (para que el
# modelo lo lea) y en una o mas "partes" con filas y columnas (para pegarlas en la hoja): las hojas de un libro, las
# tablas de un documento. Sin bibliotecas nuevas: .docx, .odt y .ods son zip con XML; .xlsx va con openpyxl.
# Los formatos viejos (.xls, .xlsb, .doc, .rtf) se abren con Excel o Word por COM, siempre en una instancia propia que
# se cierra al terminar: nunca en la del usuario.
import datetime
import os
import re
import tempfile
import zipfile
import xml.etree.ElementTree as ET

import excel
from excel import ExcelError

NL = chr(10)
TAB = chr(9)

IMAGENES = ('.png', '.jpg', '.jpeg', '.bmp', '.tif', '.tiff', '.webp', '.gif')
TEXTOS = ('.txt', '.log', '.md', '.markdown', '.csv', '.tsv', '.json', '.xml', '.html', '.htm')
PLANILLAS = ('.xlsx', '.xlsm', '.xls', '.xlsb', '.ods')
DOCUMENTOS = ('.docx', '.docm', '.odt', '.doc', '.rtf', '.pdf')
CONTACTOS = ('.vcf',)
ACEPTADOS = IMAGENES + TEXTOS + PLANILLAS + DOCUMENTOS + CONTACTOS

# Para el cuadro de dialogo del boton del clip.
TIPOS_DIALOGO = [
    ('Todos los que entiendo', ' '.join('*' + e for e in ACEPTADOS)),
    ('Planillas (Excel, OpenOffice)', '*.xlsx *.xlsm *.xls *.xlsb *.ods *.csv *.tsv'),
    ('Documentos (Word, PDF, OpenOffice, Markdown)', '*.docx *.docm *.doc *.pdf *.odt *.rtf *.md *.txt'),
    ('Contactos (agenda del telefono)', '*.vcf *.csv'),
    ('Imagenes', ' '.join('*' + e for e in IMAGENES)),
    ('Todos los archivos', '*.*'),
]


def leer(ruta):
    """Devuelve {'clase', 'texto', 'partes': [{'nombre', 'filas'}]}. Las filas de planillas conservan el tipo del
    valor (numero, fecha); las de textos son cadenas."""
    ext = os.path.splitext(ruta)[1].lower()
    if ext in IMAGENES:
        texto = excel.ocr_imagen(ruta)
        return _de_texto('imagen', texto)
    if ext in CONTACTOS:
        import contactos
        filas = contactos.filas_de_vcard(excel.leer_texto(ruta))
        return {'clase': 'contactos', 'texto': _texto_de_filas(filas), 'partes': [{'nombre': 'Contactos', 'filas': filas}]}
    if ext in ('.csv', '.tsv'):
        import contactos
        filas = contactos.filas_de_csv(excel.leer_texto(ruta))
        reducidas = contactos.reducir_contactos(filas)
        if reducidas:
            return {'clase': 'contactos', 'texto': _texto_de_filas(reducidas), 'partes': [{'nombre': 'Contactos', 'filas': reducidas}]}
        return {'clase': 'texto', 'texto': _texto_de_filas(filas), 'partes': [{'nombre': 'Datos', 'filas': filas}]}
    if ext in ('.md', '.markdown'):
        texto = excel.leer_texto(ruta)
        tablas = tablas_markdown(texto)
        if tablas:
            return {'clase': 'documento', 'texto': texto,
                    'partes': [{'nombre': 'Tabla ' + str(i), 'filas': t} for i, t in enumerate(tablas, 1)]}
        return _de_texto('texto', texto)
    if ext in ('.xlsx', '.xlsm'):
        return _planilla(leer_xlsx(ruta))
    if ext == '.ods':
        return _planilla(leer_ods(ruta))
    if ext in ('.xls', '.xlsb'):
        return _planilla(leer_con_excel(ruta))
    if ext in ('.docx', '.docm'):
        return _documento(*leer_docx(ruta))
    if ext == '.odt':
        return _documento(*leer_odt(ruta))
    if ext in ('.doc', '.rtf'):
        return _documento(*leer_con_word(ruta))
    if ext == '.pdf':
        return leer_pdf(ruta)
    if _parece_binario(ruta):
        raise ExcelError('No se leer archivos ' + (ext or 'sin extension') + '. Acepto: ' + ', '.join(ACEPTADOS))
    return _de_texto('texto', excel.leer_texto(ruta))


# ------------------------------------------------------------ armado del resultado

def _de_texto(clase, texto):
    return {'clase': clase, 'texto': texto, 'partes': [{'nombre': 'Texto', 'filas': excel.partir_filas(texto)}]}


def _planilla(hojas):
    partes = [{'nombre': n, 'filas': f} for n, f in hojas if f]
    if not partes:
        return {'clase': 'planilla', 'texto': '', 'partes': [{'nombre': 'Vacia', 'filas': []}]}
    bloques = []
    for p in partes:
        ancho = max(len(f) for f in p['filas'])
        bloques.append('== Hoja ' + p['nombre'] + ' (' + str(len(p['filas'])) + ' filas x ' + str(ancho) + ' columnas) ==' + NL
                       + _texto_de_filas(p['filas']))
    return {'clase': 'planilla', 'texto': (NL + NL).join(bloques), 'partes': partes}


def _documento(texto, tablas):
    partes = [{'nombre': 'Tabla ' + str(i), 'filas': t} for i, t in enumerate(tablas, 1) if t]
    if not partes:
        partes = [{'nombre': 'Texto', 'filas': excel.partir_filas(texto)}]
    return {'clase': 'documento', 'texto': texto, 'partes': partes}


def mostrar(v):
    'Un valor como se veria escrito: sin .0 en los enteros y las fechas como dd/mm/aaaa.'
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'VERDADERO' if v else 'FALSO'
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    if isinstance(v, datetime.datetime):
        return v.strftime('%d/%m/%Y') + ('' if v.time() == datetime.time() else v.strftime(' %H:%M'))
    if isinstance(v, datetime.date):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, datetime.time):
        return v.strftime('%H:%M')
    return str(v)


def _texto_de_filas(filas):
    return NL.join(TAB.join(mostrar(v) for v in f) for f in filas)


def _recortar(filas):
    'Saca las filas vacias del final y las celdas vacias del final de cada fila.'
    salida = []
    for f in filas:
        f = list(f)
        while f and (f[-1] is None or f[-1] == ''):
            f.pop()
        salida.append(f)
    while salida and not salida[-1]:
        salida.pop()
    return salida


def _parece_binario(ruta):
    with open(ruta, 'rb') as f:
        return b'\x00' in f.read(4096)


# ------------------------------------------------------------ PDF

_COLUMNAS_PDF = re.compile(r' {2,}')


def leer_pdf(ruta):
    """El texto de cada pagina con pypdf en modo 'layout', que respeta la posicion: las columnas de una tabla quedan
    separadas por dos o mas espacios (en el modo comun quedaban pegadas). Las paginas sin texto (un PDF escaneado) se
    leen por OCR desde la imagen que traen adentro. Las tablas son las tandas de lineas seguidas con 2 o mas columnas."""
    try:
        import pypdf
    except ImportError:
        raise ExcelError('Falta pypdf en el entorno virtual: correr ..\\.venv\\Scripts\\python.exe -m pip install -r '
                         'requirements.txt')
    nombre = os.path.basename(ruta)
    try:
        lector = pypdf.PdfReader(ruta)
        if lector.is_encrypted and not lector.decrypt(''):
            raise ExcelError(nombre + ' tiene contrasena: abrilo, copia el texto y pegamelo en el mensaje')
        paginas = list(lector.pages)
    except ExcelError:
        raise
    except Exception as exc:
        raise ExcelError('No pude leer el PDF ' + nombre + ': ' + str(exc)[:120])
    bloques, tablas, todas, por_ocr = [], [], [], 0
    for n, pagina in enumerate(paginas, 1):
        try:
            texto = pagina.extract_text(extraction_mode='layout') or ''
        except Exception:
            texto = pagina.extract_text() or ''
        if not texto.strip():
            texto = _ocr_pagina(pagina)
            por_ocr += 1 if texto.strip() else 0
        lineas = [l.rstrip() for l in texto.splitlines() if l.strip()]
        bloques.append('== Pagina ' + str(n) + ' ==' + NL + NL.join(lineas))
        tablas.extend(_tablas_de_lineas(lineas))
        todas.extend(lineas)
    texto = (NL + NL).join(bloques)
    if not todas:
        raise ExcelError(nombre + ' no tiene texto que se pueda leer (y el OCR no reconocio nada): mandame una captura '
                         'de pantalla de la pagina, o copia el texto y pegamelo en el mensaje')
    if por_ocr:
        texto = ('(' + str(por_ocr) + ' pagina(s) escaneadas, leidas por OCR: revisar numeros y nombres)' + NL + texto)
    partes = [{'nombre': 'Tabla ' + str(i), 'filas': t} for i, t in enumerate(tablas, 1)]
    if not partes:
        partes = [{'nombre': 'Texto', 'filas': [[c for c in _COLUMNAS_PDF.split(l.strip()) if c] for l in todas]}]
    return {'clase': 'documento', 'texto': texto, 'partes': partes}


def _tablas_de_lineas(lineas):
    tablas, actual = [], []
    for l in lineas:
        celdas = [c for c in _COLUMNAS_PDF.split(l.strip()) if c]
        if len(celdas) >= 2:
            actual.append(celdas)
        else:
            if len(actual) >= 2:
                tablas.append(actual)
            actual = []
    if len(actual) >= 2:
        tablas.append(actual)
    return tablas


def _ocr_pagina(pagina):
    'OCR de la imagen mas grande de una pagina (la hoja escaneada). Vacio si no trae imagenes o no hay OCR.'
    try:
        imagenes = list(pagina.images)
    except Exception:
        return ''
    if not imagenes:
        return ''
    img = max(imagenes, key=lambda i: len(i.data))
    fd, tmp = tempfile.mkstemp(suffix=os.path.splitext(img.name)[1] or '.png', prefix='excelagent-pdf-')
    try:
        with os.fdopen(fd, 'wb') as f:
            f.write(img.data)
        return excel.ocr_imagen(tmp)
    except ExcelError:
        return ''
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass


# ------------------------------------------------------------ Markdown

_SEPARADOR_MD = re.compile(r'^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$')


def tablas_markdown(texto):
    'Las tablas de un Markdown (| a | b |), sin la linea de guiones.'
    tablas, actual = [], []
    for linea in texto.splitlines():
        if linea.strip().startswith('|'):
            if not _SEPARADOR_MD.match(linea):
                celdas = linea.strip().strip('|').split('|')
                actual.append([c.strip() for c in celdas])
        elif actual:
            tablas.append(actual)
            actual = []
    if actual:
        tablas.append(actual)
    return [t for t in tablas if len(t) > 1]


# ------------------------------------------------------------ planillas

def leer_xlsx(ruta):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelError('No pude leer el libro ' + os.path.basename(ruta) + ': ' + str(exc)[:120])
    try:
        return [(ws.title, _recortar(ws.iter_rows(values_only=True))) for ws in wb.worksheets]
    finally:
        wb.close()


_ODS = {'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
        'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
        'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}


def _q(prefijo, nombre):
    return '{' + _ODS[prefijo] + '}' + nombre


def _texto_odf(elem):
    # text:s son espacios repetidos, text:tab y text:line-break tambien cuentan
    partes = [elem.text or '']
    for hijo in elem:
        etiqueta = hijo.tag
        if etiqueta == _q('text', 's'):
            partes.append(' ' * int(hijo.get(_q('text', 'c'), '1')))
        elif etiqueta == _q('text', 'tab'):
            partes.append(TAB)
        elif etiqueta == _q('text', 'line-break'):
            partes.append(NL)
        else:
            partes.append(_texto_odf(hijo))
        partes.append(hijo.tail or '')
    return ''.join(partes)


def _valor_ods(celda):
    tipo = celda.get(_q('office', 'value-type'))
    if tipo in ('float', 'percentage', 'currency'):
        v = float(celda.get(_q('office', 'value')))
        return int(v) if v.is_integer() and abs(v) < 1e15 else v
    if tipo == 'date':
        d = celda.get(_q('office', 'date-value'))
        try:
            return datetime.datetime.fromisoformat(d) if 'T' in d else datetime.date.fromisoformat(d)
        except ValueError:
            return d
    if tipo == 'boolean':
        return celda.get(_q('office', 'boolean-value')) == 'true'
    return NL.join(_texto_odf(p) for p in celda.findall(_q('text', 'p')))


def leer_ods(ruta):
    try:
        with zipfile.ZipFile(ruta) as z:
            raiz = ET.fromstring(z.read('content.xml'))
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        raise ExcelError('No pude leer ' + os.path.basename(ruta) + ': ' + str(exc)[:120])
    hojas = []
    for tabla in raiz.iter(_q('table', 'table')):
        filas, vacias = [], 0
        for fila in tabla.iter(_q('table', 'table-row')):
            celdas, pendientes = [], 0
            for c in fila:
                if c.tag not in (_q('table', 'table-cell'), _q('table', 'covered-table-cell')):
                    continue
                n = int(c.get(_q('table', 'number-columns-repeated'), '1'))
                v = _valor_ods(c)
                if v == '' or v is None:
                    pendientes += n          # las vacias repetidas (1024 al final) solo cuentan si sigue algo
                    continue
                celdas += [''] * pendientes + [v] * min(n, 1024)
                pendientes = 0
            veces = int(fila.get(_q('table', 'number-rows-repeated'), '1'))
            if not celdas:
                vacias += veces              # idem con las filas vacias repetidas del final
                continue
            filas += [[]] * min(vacias, 10000) + [celdas] * min(veces, 10000)
            vacias = 0
        hojas.append((tabla.get(_q('table', 'name'), 'Hoja'), _recortar(filas)))
    return hojas


def _sin_zona(v):
    if isinstance(v, datetime.datetime):           # pywintypes llega con zona horaria: se deja la hora local tal cual
        return datetime.datetime(v.year, v.month, v.day, v.hour, v.minute, v.second)
    if type(v).__name__ == 'Decimal':              # moneda
        return float(v)
    return v


def leer_con_excel(ruta):
    'Formatos que solo Excel lee (.xls, .xlsb): en una instancia propia, sin ventana, solo lectura.'
    import pythoncom
    import win32com.client as win32
    pythoncom.CoInitialize()
    app = win32.DispatchEx('Excel.Application')
    try:
        app.DisplayAlerts = False
        app.Visible = False
        wb = app.Workbooks.Open(os.path.abspath(ruta), 0, True)      # sin actualizar vinculos, solo lectura
        try:
            hojas = []
            for ws in wb.Worksheets:
                v = ws.UsedRange.Value
                if not isinstance(v, tuple):
                    v = ((v,),)
                hojas.append((str(ws.Name), _recortar([[_sin_zona(x) for x in f] for f in v])))
            return hojas
        finally:
            wb.Close(False)
    except ExcelError:
        raise
    except Exception as exc:
        raise ExcelError('Excel no pudo leer ' + os.path.basename(ruta) + ': ' + str(exc)[:120])
    finally:
        app.Quit()


# ------------------------------------------------------------ documentos

_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _texto_w(elem):
    partes = []
    for e in elem.iter():
        if e.tag == _W + 't':
            partes.append(e.text or '')
        elif e.tag == _W + 'tab':
            partes.append(TAB)
        elif e.tag in (_W + 'br', _W + 'cr'):
            partes.append(NL)
    return ''.join(partes)


def leer_docx(ruta):
    'Texto y tablas de un .docx, en el orden del documento.'
    try:
        with zipfile.ZipFile(ruta) as z:
            raiz = ET.fromstring(z.read('word/document.xml'))
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        raise ExcelError('No pude leer ' + os.path.basename(ruta) + ': ' + str(exc)[:120])
    cuerpo = raiz.find(_W + 'body')
    lineas, tablas = [], []
    for hijo in (cuerpo if cuerpo is not None else []):
        if hijo.tag == _W + 'p':
            lineas.append(_texto_w(hijo))
        elif hijo.tag == _W + 'tbl':
            filas = [[' '.join(_texto_w(p) for p in tc.iter(_W + 'p')).strip() for tc in tr.findall(_W + 'tc')]
                     for tr in hijo.findall(_W + 'tr')]
            tablas.append(_recortar(filas))
            lineas += [TAB.join(f) for f in filas]
    return NL.join(lineas).strip(), tablas


def leer_odt(ruta):
    try:
        with zipfile.ZipFile(ruta) as z:
            raiz = ET.fromstring(z.read('content.xml'))
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        raise ExcelError('No pude leer ' + os.path.basename(ruta) + ': ' + str(exc)[:120])
    lineas, tablas = [], []

    def recorrer(elem):
        for hijo in elem:
            if hijo.tag in (_q('text', 'p'), _q('text', 'h')):
                lineas.append(_texto_odf(hijo))
            elif hijo.tag == _q('table', 'table'):
                filas = [[' '.join(_texto_odf(p) for p in c.iter(_q('text', 'p'))).strip()
                          for c in fila if c.tag == _q('table', 'table-cell')]
                         for fila in hijo.iter(_q('table', 'table-row'))]
                tablas.append(_recortar(filas))
                lineas.extend(TAB.join(f) for f in filas)
            else:
                recorrer(hijo)                   # listas, secciones
    texto_oficina = raiz.find(_q('office', 'body'))
    recorrer(texto_oficina if texto_oficina is not None else raiz)
    return NL.join(lineas).strip(), tablas


def leer_con_word(ruta):
    """.doc y .rtf: Word (instancia propia, sin ventana) lo guarda como .docx en una carpeta temporal y se lee como
    cualquier .docx. Sin Word instalado no hay forma: se dice como salir del paso."""
    import pythoncom
    import win32com.client as win32
    pythoncom.CoInitialize()
    try:
        word = win32.DispatchEx('Word.Application')
    except Exception:
        raise ExcelError('Para leer ' + os.path.basename(ruta) + ' hace falta Microsoft Word, y no esta instalado. '
                         'Abrilo con otro programa (por ejemplo LibreOffice) y guardalo como .docx u .odt, o copia el '
                         'texto y pegalo en el mensaje.')
    carpeta = tempfile.mkdtemp(prefix='excelagent_doc_')
    destino = os.path.join(carpeta, 'copia.docx')
    try:
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(os.path.abspath(ruta), False, True, False)   # sin convertir, solo lectura
        try:
            doc.SaveAs2(destino, 16)                                             # wdFormatDocumentDefault (.docx)
        finally:
            doc.Close(False)
        return leer_docx(destino)
    except ExcelError:
        raise
    except Exception as exc:
        raise ExcelError('Word no pudo leer ' + os.path.basename(ruta) + ': ' + str(exc)[:120])
    finally:
        word.Quit()
        try:
            os.remove(destino)
        except OSError:
            pass
        try:
            os.rmdir(carpeta)
        except OSError:
            pass
